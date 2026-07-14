#!/usr/bin/env python3
"""Editable flyers: one .pptx, 3 slides at 1080x1350, every element a layer."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from pptx.oxml.ns import qn
from PIL import Image
import copy

PX = 9525  # EMU per pixel at 96dpi
W, H = 1080 * PX, 1350 * PX
IMG = "/home/user/content-system/main-site"
OUT = "/home/user/content-system/brand/flyers/flyers-editable.pptx"

ELECTRIC = RGBColor(0x2C, 0x4B, 0xE0)
ACCENT_LT = RGBColor(0x8F, 0xA6, 0xFF)
DARK = RGBColor(0x0C, 0x0C, 0x10)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREY = RGBColor(0xA9, 0xA9, 0xB4)
OFFWHITE = RGBColor(0xED, 0xED, 0xF2)

FLYERS = [
    dict(photo=f"{IMG}/Portrait.jpeg", focus=(0.5, 0.2),
         kicker="LIVE MASTERCLASS · 90 MINUTES",
         h1=[("Hire your first", WHITE, False), ("AI employee.", ACCENT_LT, True)],
         bullets=["It is running before we hang up",
                  "Live, step by step, with me",
                  "No code. No jargon. No homework."],
         cta="SEATS LIMITED · LINK IN BIO"),
    dict(photo=f"{IMG}/fatiha-desk.jpg", focus=(0.72, 0.5),
         kicker="LIVE BOOTCAMP · 2 DAYS",
         h1=[("Your business,", WHITE, False), ("running on AI.", WHITE, False), ("In 2 days.", ACCENT_LT, True)],
         bullets=["Day 1: your AI sounds like you, inbox handled",
                  "Day 2: a month of content, one system live",
                  "Replays and every template included"],
         cta="THE WAITLIST HEARS THE DATE FIRST · LINK IN BIO"),
    dict(photo=f"{IMG}/fatiha-studio.jpg", focus=(0.5, 0.3),
         kicker="GET IT DONE DAY · MAX 10 SEATS",
         h1=[("The Get It Done", WHITE, False), ("Day.", ACCENT_LT, True)],
         bullets=["Come with the half-finished automation",
                  "Build it with me beside you",
                  "It runs before dinner. Ten seats."],
         cta="TEN SEATS ONLY · LINK IN BIO"),
]

def set_alpha(shape, pct):
    """Set fill transparency (pct visible alpha, e.g. 75 = 75% opaque)."""
    fill = shape.fill._xPr.find(qn('a:solidFill'))
    clr = fill.find(qn('a:srgbClr'))
    a = clr.makeelement(qn('a:alpha'), {'val': str(int(pct * 1000))})
    clr.append(a)

def rect(slide, x, y, w, h, color, line=False):
    s = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    s.fill.solid()
    s.fill.fore_color.rgb = color
    s.line.fill.background()
    s.shadow.inherit = False
    return s

def textbox(slide, x, y, w, h, runs_lines, font, size, bold=False,
            spacing=None, align=PP_ALIGN.LEFT, letterspace=None):
    tb = slide.shapes.add_textbox(Emu(x), Emu(y), Emu(w), Emu(h))
    tf = tb.text_frame
    tf.word_wrap = True
    for i, runs in enumerate(runs_lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = align
        if spacing:
            p.line_spacing = spacing
        for text, color, italic in runs:
            r = p.add_run()
            r.text = text
            r.font.name = font
            r.font.size = Pt(size)
            r.font.bold = bold
            r.font.italic = italic
            r.font.color.rgb = color
            if letterspace:
                r.font._rPr.set('spc', str(letterspace))
    return tb

prs = Presentation()
prs.slide_width = Emu(W)
prs.slide_height = Emu(H)
blank = prs.slide_layouts[6]

for f in FLYERS:
    slide = prs.slides.add_slide(blank)

    # background
    rect(slide, 0, 0, W, H, DARK)

    # photo, cover-cropped
    im = Image.open(f["photo"])
    img_as, slide_as = im.width / im.height, 1080 / 1350
    pic = slide.shapes.add_picture(f["photo"], 0, 0, Emu(W), Emu(H))
    fx, fy = f["focus"]
    if img_as > slide_as:
        total = 1 - slide_as / img_as
        pic.crop_left, pic.crop_right = fx * total, (1 - fx) * total
    else:
        total = 1 - img_as / slide_as
        pic.crop_top, pic.crop_bottom = fy * total, (1 - fy) * total

    # dark shade over the lower half (two editable layers standing in for the gradient)
    set_alpha(rect(slide, 0, int(H * 0.42), W, int(H * 0.30), DARK), 55)
    set_alpha(rect(slide, 0, int(H * 0.72), W, int(H * 0.28), DARK), 88)

    # top line
    rect(slide, 0, 0, W, 16 * PX, ELECTRIC)

    # kicker chip
    chip_w = (44 + int(len(f["kicker"]) * 14.5)) * PX
    chip = rect(slide, 64 * PX, 56 * PX, chip_w, 54 * PX, ELECTRIC)
    ctf = chip.text_frame
    ctf.word_wrap = False
    p = ctf.paragraphs[0]
    p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f["kicker"]
    r.font.name = "Space Mono"; r.font.size = Pt(17); r.font.bold = True
    r.font.color.rgb = WHITE

    # headline
    n = len(f["h1"])
    h1_y = (1350 - 150 - 96 - 340 - n * 116) * PX
    textbox(slide, 64 * PX, h1_y, 952 * PX, (n * 116 + 20) * PX,
            [[(t, c, i)] for t, c, i in f["h1"]],
            "Playfair Display", 78, spacing=1.02)

    # rule
    rule_y = h1_y + (n * 116 + 30) * PX
    rect(slide, 64 * PX, rule_y, 220 * PX, 7 * PX, ELECTRIC)

    # bullets
    by = rule_y + 34 * PX
    lines = [[("→  ", ACCENT_LT, False), (b, OFFWHITE, False)] for b in f["bullets"]]
    textbox(slide, 64 * PX, by, 952 * PX, 170 * PX, lines, "Inter", 24, spacing=1.4)

    # who line
    who_y = by + 180 * PX
    textbox(slide, 64 * PX, who_y, 952 * PX, 70 * PX,
            [[("FATIHA CHIKH", WHITE, False),
              ("  ·  SHIFT & LEAD  ·  RUNS HER BUSINESS ON AI EMPLOYEES SHE BUILT HERSELF", GREY, False)]],
            "Space Mono", 15)

    # CTA bar
    bar = rect(slide, 0, (1350 - 96) * PX, W, 96 * PX, ELECTRIC)
    btf = bar.text_frame
    p = btf.paragraphs[0]; p.alignment = PP_ALIGN.CENTER
    r = p.add_run(); r.text = f["cta"]
    r.font.name = "Inter"; r.font.size = Pt(22); r.font.bold = True
    r.font.color.rgb = WHITE

prs.save(OUT)
import os
print("saved", OUT, os.path.getsize(OUT) // 1024, "KB")
